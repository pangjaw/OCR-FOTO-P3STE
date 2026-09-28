#!/usr/bin/env python3
"""Ekspor nama file dan daftar aset dalam setiap PDF ke Excel."""
import argparse
import json
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def _asset_code(funcloc: str) -> str:
    """Ambil kode sebelum titik dua, contoh AXL11540."""
    return funcloc.split(":", 1)[0].strip() if ":" in funcloc else "-"


def _pdf_asset_codes(pdf_path: Path, category: str) -> dict[str, str]:
    """Fallback untuk schedule lama yang belum memiliki field funcloc."""
    try:
        import fitz
        from export_pdf_foto import extract_all_funclocs, extract_identifier
        with fitz.open(pdf_path) as doc:
            text = doc[0].get_text() if len(doc) else ""
        return {
            extract_identifier(line, category): _asset_code(line)
            for line in extract_all_funclocs(text)
            if extract_identifier(line, category)
        }
    except Exception:
        return {}


def export_assets(schedule_path: Path, pdf_dir: Path, output_path: Path) -> tuple[int, int]:
    schedules = json.loads(schedule_path.read_text(encoding="utf-8")).get("schedules", [])
    assets_by_path = defaultdict(list)
    assets_by_name = defaultdict(list)
    dates_by_path = {}
    dates_by_name = {}
    for item in schedules:
        identifier = item.get("identifier") or item.get("pdf_stem") or ""
        iso_date = item.get("iso_date", "")
        date_text = f"{iso_date[8:10]}-{iso_date[5:7]}-{iso_date[:4]}" if len(iso_date) == 10 else "-"
        asset = {
            "identifier": identifier,
            "code": _asset_code(item.get("funcloc", "")),
            "category": item.get("category", ""),
        }
        raw_path = item.get("pdf_path")
        if raw_path:
            path_key = str(Path(raw_path).resolve())
            dates_by_path[path_key] = date_text
            if identifier:
                assets_by_path[path_key].append(asset)
        name_key = item.get("file", "")
        dates_by_name[name_key] = date_text
        if identifier:
            assets_by_name[name_key].append(asset)

    pdf_files = sorted(pdf_dir.rglob("*.pdf"))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Aset"
    ws.append(["TANGGAL", "FILE NAME", "KODE ASET", "ASET"])
    asset_count = 0

    for path in pdf_files:
        path_key = str(path.resolve())
        values = assets_by_path.get(path_key) or assets_by_name.get(path.name, [])
        date_text = dates_by_path.get(path_key) or dates_by_name.get(path.name, "-")
        unique = list({(v["code"], v["identifier"]): v for v in values}.values())
        if unique and any(v["code"] == "-" for v in unique):
            fallback = _pdf_asset_codes(path, unique[0]["category"])
            for value in unique:
                if value["code"] == "-":
                    value["code"] = fallback.get(value["identifier"], "-")
        if unique:
            for value in unique:
                ws.append([date_text, path.name, value["code"], value["identifier"]])
                asset_count += 1
        else:
            ws.append([date_text, path.name, "-", "-"])

    fill = PatternFill("solid", fgColor="1E3A8A")
    font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row[0].row].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:D1"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 40
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return len(pdf_files), asset_count


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--schedule", required=True)
    parser.add_argument("--pdf-dir", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    files, assets = export_assets(Path(args.schedule), Path(args.pdf_dir), Path(args.output))
    print(json.dumps({"success": True, "files": files, "assets": assets, "output": args.output}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
