"""Export schedule.json ke Excel ringkasan audit."""
import json
import sys
from pathlib import Path
from datetime import datetime

APP_DIR = Path(__file__).resolve().parent
SCHEDULE_PATH = APP_DIR / "schedule.json"

with open(SCHEDULE_PATH, "r", encoding="utf-8") as f:
    data = json.load(f)

rows = []
for s in data.get("schedules", []):
    photos = s.get("photos", {})
    for photo_name, iso_time in photos.items():
        # Parse timestamp ke jam menit
        dt = datetime.fromisoformat(iso_time)
        jam_str = dt.strftime("%H:%M")
        rows.append({
            "Tim": s.get("tim", ""),
            "BTP": s.get("btp", ""),
            "Kategori": s.get("category", ""),
            "Identifier": s.get("identifier", ""),
            "Tanggal": s.get("date", ""),
            "Foto": photo_name,
            "Jam": jam_str,
            "Waktu (menit)": s.get("waktu_menit", ""),
            "File PDF": s.get("file", ""),
        })

# ── Excel ──
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl not installed, using CSV fallback")
    import csv
    csv_path = APP_DIR / "schedule_export.csv"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=rows[0].keys())
        w.writeheader()
        w.writerows(rows)
    print(f"CSV: {csv_path} ({len(rows)} baris)")
    sys.exit(0)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Schedule Audit"

# Headers
headers = list(rows[0].keys())
header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)

for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

# Zebra rows
even_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

for row_idx, row in enumerate(rows, 2):
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=row[h])
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center")
        if row_idx % 2 == 0:
            cell.fill = even_fill

# Auto-width
for col_idx in range(1, len(headers) + 1):
    max_len = max(
        len(str(ws.cell(row=row_idx, column=col_idx).value or ""))
        for row_idx in range(1, len(rows) + 2)
    )
    ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

# Freeze header
ws.freeze_panes = "A2"

# Summary sheet
ws2 = wb.create_sheet("Ringkasan")
summary = {}
for r in rows:
    key = (r["Tim"], r["BTP"], r["Kategori"])
    summary[key] = summary.get(key, 0) + 1

ws2.append(["Tim", "BTP", "Kategori", "Jumlah Foto"])
for (tim, btp, cat), count in sorted(summary.items()):
    ws2.append([tim, btp, cat, count])

# Total row
ws2.append([])
ws2.append(["", "", "TOTAL", sum(summary.values())])

for row in ws2.iter_rows(min_row=1, max_row=1):
    for cell in row:
        cell.fill = header_fill
        cell.font = header_font

output_path = APP_DIR / "schedule_audit.xlsx"
wb.save(output_path)
print(f"Excel: {output_path}")
print(f"Total: {len(rows)} baris, {len(data['schedules'])} aset")

# Quick validation
tim_counts = {}
for s in data["schedules"]:
    t = s.get("tim", 0)
    tim_counts[t] = tim_counts.get(t, 0) + 1
print(f"Distribusi Tim: {dict(sorted(tim_counts.items()))}")
